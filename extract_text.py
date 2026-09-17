import os
import json
from docx import Document

INPUT_FOLDER = "policies 2"
OUTPUT_FOLDER = "extracted_structure"

os.makedirs(OUTPUT_FOLDER, exist_ok=True)


def heading_level(style_name):
    """Return an integer heading level (1, 2, 3...) or None if not a heading."""
    style_name = (style_name or "").lower()
    if style_name == "title":
        return 0
    if style_name.startswith("heading"):
        parts = style_name.split()
        if len(parts) == 2 and parts[1].isdigit():
            return int(parts[1])
    return None


def extract_table(table):
    """Convert a docx table into a list of row lists (plain text cells)."""
    rows = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows.append(cells)
    return rows


def extract_structure(file_path):
    """
    Walk the document top-to-bottom and produce an ordered list of elements:
      {"type": "heading", "level": int, "text": str}
      {"type": "paragraph", "text": str}
      {"type": "table", "rows": [[...], ...]}
    This preserves document order and hierarchy instead of flattening
    everything into one blob of text.
    """
    doc = Document(file_path)
    elements = []

    # python-docx doesn't give a single ordered stream of paragraphs+tables
    # by default, so we walk the underlying XML body in document order.
    body = doc.element.body
    para_map = {p._p: p for p in doc.paragraphs}
    table_map = {t._tbl: t for t in doc.tables}

    for child in body.iterchildren():
        if child in para_map:
            para = para_map[child]
            text = para.text.strip()
            if not text:
                continue
            level = heading_level(para.style.name if para.style else None)
            if level is not None:
                elements.append({"type": "heading", "level": level, "text": text})
            else:
                elements.append({"type": "paragraph", "text": text})
        elif child in table_map:
            table = table_map[child]
            rows = extract_table(table)
            if rows:
                elements.append({"type": "table", "rows": rows})

    return elements


def main():
    if not os.path.exists(INPUT_FOLDER):
        print(f"Folder '{INPUT_FOLDER}' not found. Place it next to this script.")
        return

    docx_files = [f for f in os.listdir(INPUT_FOLDER) if f.lower().endswith(".docx")]

    if not docx_files:
        print(f"No .docx files found in '{INPUT_FOLDER}'.")
        return

    for filename in docx_files:
        file_path = os.path.join(INPUT_FOLDER, filename)
        elements = extract_structure(file_path)

        doc_title = next(
            (e["text"] for e in elements if e["type"] == "heading" and e["level"] == 0),
            os.path.splitext(filename)[0]
        )

        out_data = {
            "source": os.path.splitext(filename)[0],
            "title": doc_title,
            "elements": elements
        }

        out_name = os.path.splitext(filename)[0] + ".json"
        out_path = os.path.join(OUTPUT_FOLDER, out_name)

        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(out_data, f, indent=2, ensure_ascii=False)

        n_headings = sum(1 for e in elements if e["type"] == "heading")
        n_tables = sum(1 for e in elements if e["type"] == "table")
        print(f"Extracted: {filename} -> {out_path} "
              f"({len(elements)} elements, {n_headings} headings, {n_tables} tables)")


if __name__ == "__main__":
    main()